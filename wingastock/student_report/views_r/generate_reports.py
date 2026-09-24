import json
import re
from io import BytesIO
from urllib.request import Request, urlopen

from PIL import Image
import cloudinary.uploader

from bs4 import BeautifulSoup

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from student_report.models import (
    FieldPlacement,
    DailyFieldReport,
    GeneratedFieldReport,
)


# =========================================================
# GENERATE FIELD REPORT PAGE
# =========================================================
@login_required(login_url="home")
def generate_field_report(request, id):
    placement = get_object_or_404(FieldPlacement, id=id)

    if placement.student != request.user:
        messages.error(request, "You are not authorized to access this field report.")
        return redirect("daily_reports")

    reports = placement.daily_reports.all().order_by("date", "id")

    saved_report = (
        GeneratedFieldReport.objects
        .filter(placement=placement)
        .first()
    )

    return render(
        request,
        "generate_field_re.html",
        {
            "placement": placement,
            "reports": reports,
            "saved_report": saved_report,
        }
    )


# =========================================================
# SAVE GENERATED FIELD REPORT DATA
# =========================================================
@login_required(login_url="home")
@require_POST
def save_field_report(request, id):

    placement = get_object_or_404(
        FieldPlacement,
        id=id,
        student=request.user
    )

    report, created = GeneratedFieldReport.objects.get_or_create(
        placement=placement
    )

    text_fields = [
        "report_html",
        "declaration",
        "abbreviations",
        "historical_background",
        "organization_structure",
        "vision",
        "mission",
        "objectives",
        "section_21_brief_overview",
        "section_22_activities_performed",
        "section_23_general_observation",
        "section_24_challenges",
        "section_25_solutions",
        "conclusion",
        "recommendations",
        "references_bibliography",
    ]

    changed_fields = []

    for field in text_fields:
        if field in request.POST:
            value = request.POST.get(field, "")
            setattr(report, field, value)
            changed_fields.append(field)

    # -------------------------------------------------
    # IMAGE UPLOAD
    # -------------------------------------------------
    image_field = request.POST.get("image_field", "").strip()
    uploaded_image = request.FILES.get("image")

    allowed_image_fields = {
        "historical_background_image",
        "organization_structure_image",
        "organization_image",
        "activity_image",
        "appendix_image",
    }

    if uploaded_image:

        if image_field not in allowed_image_fields:
            return JsonResponse(
                {"success": False, "message": "Invalid report image field."},
                status=400
            )

        if uploaded_image.size > 5 * 1024 * 1024:
            return JsonResponse(
                {"success": False, "message": "Image must not be larger than 5 MB."},
                status=400
            )

        allowed_types = {
            "image/jpeg",
            "image/png",
            "image/webp",
            "image/gif",
        }

        if uploaded_image.content_type not in allowed_types:
            return JsonResponse(
                {
                    "success": False,
                    "message": "Only JPG, PNG, WEBP and GIF images are allowed.",
                },
                status=400
            )

        try:
            image = Image.open(uploaded_image)
            image.verify()
            uploaded_image.seek(0)
        except Exception:
            return JsonResponse(
                {"success": False, "message": "Invalid image file."},
                status=400
            )

        old_image = getattr(report, image_field, None)
        old_public_id = None

        if old_image:
            try:
                old_public_id = old_image.public_id
            except Exception:
                old_public_id = None

        try:
            upload_result = cloudinary.uploader.upload(
                uploaded_image,
                folder="field_reports"
            )

            public_id = upload_result.get("public_id")
            secure_url = upload_result.get("secure_url")

            if not public_id:
                raise Exception("Cloudinary upload did not return a public ID.")

            setattr(report, image_field, public_id)
            report.save()

        except Exception as error:
            return JsonResponse(
                {"success": False, "message": f"Image upload failed: {error}"},
                status=500
            )

        if old_public_id:
            try:
                cloudinary.uploader.destroy(
                    old_public_id,
                    resource_type="image"
                )
            except Exception as error:
                print("OLD CLOUDINARY IMAGE DELETE ERROR:", error)

        return JsonResponse({
            "success": True,
            "message": "Report image uploaded successfully.",
            "placement_id": placement.id,
            "student_id": request.user.id,
            "report_id": report.id,
            "image_field": image_field,
            "image_url": secure_url,
        })

    if changed_fields:
        report.save()
    else:
        return JsonResponse(
            {"success": False, "message": "No report information was submitted."},
            status=400
        )

    return JsonResponse({
        "success": True,
        "message": "Report saved successfully.",
        "placement_id": placement.id,
        "student_id": request.user.id,
        "report_id": report.id,
        "saved_fields": changed_fields,
    })


# =========================================================
# REFRESH FIELD REPORT DATA
# =========================================================
@login_required(login_url="home")
@require_POST
def refresh_field_report_data(request, id):

    placement = get_object_or_404(FieldPlacement, id=id)

    if placement.student != request.user:
        return JsonResponse(
            {
                "success": False,
                "message": "You are not authorized to access this report."
            },
            status=403
        )

    reports = placement.daily_reports.all().order_by("date", "id")

    report_data = []

    for report in reports:
        report_data.append({
            "id": report.id,
            "date": report.date.strftime("%Y-%m-%d") if report.date else "",
            "title": report.title or "",
            "activities": report.activities or "",
            "skills_learned": report.skills_learned or "",
            "challenges": report.challenges or "",
            "solutions": report.solutions or "",
            "remarks": report.remarks or "",
            "supervisor_comment": report.supervisor_comment or "",
            "supervisor_approved": report.supervisor_approved,
        })

    return JsonResponse({
        "success": True,
        "reports": report_data,
        "count": len(report_data),
    })


# =========================================================
# SAVE REPORT SUMMARY FROM GPT
# =========================================================
@login_required(login_url="home")
@require_POST
def save_report_summary(request, id):

    try:
        placement = get_object_or_404(FieldPlacement, id=id)

        if placement.student != request.user:
            return JsonResponse(
                {
                    "success": False,
                    "message": "You are not authorized to save this report."
                },
                status=403
            )

        data = json.loads(request.body)
        summary = (data.get("summary", "")).strip()

        if not summary:
            return JsonResponse(
                {
                    "success": False,
                    "message": "Please paste the GPT summary first."
                },
                status=400
            )

        generated_report, created = GeneratedFieldReport.objects.get_or_create(
            placement=placement
        )

        generated_report.summary_html = summary
        generated_report.save(update_fields=["summary_html", "updated_at"])

        return JsonResponse({
            "success": True,
            "message": "Report summary saved successfully.",
            "summary": generated_report.summary_html,
            "created": created,
        })

    except json.JSONDecodeError:
        return JsonResponse(
            {
                "success": False,
                "message": "Invalid JSON data."
            },
            status=400
        )

    except Exception as e:
        return JsonResponse(
            {
                "success": False,
                "message": str(e)
            },
            status=500
        )




# ============================================================
# COMPLETENESS CHECK
# ============================================================

REQUIRED_REPORT_FIELDS = [
    ("historical_background", "1.1 Background of the Industry/Organization"),
    ("vision", "1.3.1 Vision"),
    ("mission", "1.3.2 Mission"),
    ("objectives", "1.3.3 Objectives"),
    ("section_21_brief_overview", "2.1 Brief Overview of the Undertaking"),
    ("section_22_activities_performed", "2.2 Activity/Activities Performed During IPT"),
    ("section_23_general_observation", "2.3 General Observations"),
    ("section_24_challenges", "2.4 Challenges/Problems Faced"),
    ("section_25_solutions", "2.5 How the Student Solved the Challenges"),
    ("conclusion", "3.1 Conclusion"),
    ("recommendations", "3.2 Recommendations"),
    ("references_bibliography", "References"),
]


def get_missing_report_fields(generated_report):
    if generated_report is None:
        return list(REQUIRED_REPORT_FIELDS)

    missing = []

    for field, label in REQUIRED_REPORT_FIELDS:
        value = getattr(generated_report, field, "") or ""

        if not str(value).strip():
            missing.append((field, label))

    return missing


# =========================================================
# HELPER: CELL TEXT / SHADING / WIDTH
# =========================================================
def set_cell_text(cell, text, bold=False):
    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        str(text) if text is not None else ""
    )

    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill="E7E6E6"):
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def set_cell_width(cell, width_mm):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_w = tc_pr.first_child_found_in("w:tcW")

    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)

    tc_w.set(qn("w:w"), str(int(width_mm * 56.7)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")

    tr_pr.append(tbl_header)


# =========================================================
# HELPER: PARAGRAPH FORMAT
# =========================================================
def set_paragraph_format(paragraph):
    paragraph.paragraph_format.line_spacing = 1.6
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Mm(10)

    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


# =========================================================
# HELPER: BODY PARAGRAPH
# =========================================================
def add_body_paragraph(document, text):
    text = str(text or "").strip()

    if not text:
        return

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    run = paragraph.add_run(text)

    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    set_paragraph_format(paragraph)

    return paragraph


# =========================================================
# HELPER: NORMAL PARAGRAPH
# =========================================================
def add_normal_paragraph(document, text, bold=False):
    text = str(text or "").strip()

    if not text:
        return

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = None

    run = paragraph.add_run(text)

    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    return paragraph


# =========================================================
# HELPER: HEADING 1
# =========================================================
def add_heading_1(document, text):
    paragraph = document.add_paragraph()

    try:
        paragraph.style = document.styles["Heading 1"]
    except KeyError:
        pass

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = Mm(0)

    run = paragraph.add_run(str(text).upper())

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    return paragraph


# =========================================================
# HELPER: HEADING 2
# =========================================================
def add_heading_2(document, text):
    paragraph = document.add_paragraph()

    try:
        paragraph.style = document.styles["Heading 2"]
    except KeyError:
        pass

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = Mm(0)

    run = paragraph.add_run(str(text))

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    return paragraph


# =========================================================
# HELPER: HEADING 3
# =========================================================
def add_heading_3(document, text):
    paragraph = document.add_paragraph()

    try:
        paragraph.style = document.styles["Heading 3"]
    except KeyError:
        pass

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = Mm(6)

    run = paragraph.add_run(str(text))

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    return paragraph


# =========================================================
# HELPER: CENTER HEADING
# =========================================================
def add_center_heading(document, text, size=14, bold=True):
    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.3

    run = paragraph.add_run(str(text))

    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)

    return paragraph


# =========================================================
# HELPER: DETAIL LINE FOR COVER
# =========================================================
def add_detail_line(document, label, value):
    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.left_indent = Mm(18)
    paragraph.paragraph_format.right_indent = Mm(18)

    run_label = paragraph.add_run(str(label) + " ")

    run_label.bold = True
    run_label.font.name = "Times New Roman"
    run_label.font.size = Pt(12)

    run_value = paragraph.add_run(
        str(value) if value else "-"
    )

    run_value.font.name = "Times New Roman"
    run_value.font.size = Pt(12)

    return paragraph


# =========================================================
# HELPER: PAGE BREAK
# =========================================================
def add_page_break(document):
    document.add_page_break()


# =========================================================
# HELPER: PAGE NUMBER
# =========================================================
def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None

    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


# =========================================================
# HELPER: PAGE BORDER
# =========================================================
def add_page_border(section, sz=24, space=24, color="000000"):
    sect_pr = section._sectPr

    for existing in sect_pr.findall(qn("w:pgBorders")):
        sect_pr.remove(existing)

    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")

    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")

        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(sz))
        border.set(qn("w:space"), str(space))
        border.set(qn("w:color"), color)

        pg_borders.append(border)

    sect_pr.append(pg_borders)


def remove_page_border(section):
    sect_pr = section._sectPr

    for existing in sect_pr.findall(qn("w:pgBorders")):
        sect_pr.remove(existing)


# =========================================================
# HELPER: CONFIGURE SECTION
# =========================================================
def configure_section(section, with_border=False):
    section.page_width = Mm(210)
    section.page_height = Mm(297)

    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    if with_border:
        add_page_border(section)
    else:
        remove_page_border(section)


# =========================================================
# HELPER: TOC FIELD
# =========================================================
def add_toc_field(document):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_after = Pt(6)

    run_begin = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    run_begin._r.append(fld_begin)

    run_instr = paragraph.add_run()

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    run_instr._r.append(instr)

    run_sep = paragraph.add_run()

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    run_sep._r.append(fld_sep)

    run_placeholder = paragraph.add_run(
        'Right-click here and select "Update Field" to generate the Table of Contents.'
    )

    run_placeholder.italic = True
    run_placeholder.font.name = "Times New Roman"
    run_placeholder.font.size = Pt(11)
    run_placeholder.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    run_end = paragraph.add_run()

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run_end._r.append(fld_end)

    return paragraph


def enable_update_fields(document):
    settings = document.settings.element

    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")

    settings.append(update_fields)


# =========================================================
# HELPER: CLEAN TEXT
# =========================================================
def clean_text(text):
    if not text:
        return ""

    text = str(text).strip()

    if "<" in text and ">" in text:
        soup = BeautifulSoup(text, "html.parser")

        for tag in soup.find_all(
            [
                "p",
                "div",
                "br",
                "li",
                "h1",
                "h2",
                "h3",
                "h4",
                "h5",
                "h6",
                "ul",
                "ol",
            ]
        ):
            if tag.name == "br":
                tag.replace_with("\n")

            elif tag.name == "li":
                item_text = tag.get_text(" ", strip=True)
                tag.replace_with(f"\n{item_text}")

            elif tag.name in ("ul", "ol"):
                continue

            else:
                tag_text = tag.get_text(" ", strip=True)
                tag.replace_with(f"\n{tag_text}\n")

        text = soup.get_text("\n", strip=False)

    text = re.sub(r"```[a-zA-Z0-9]*\n?", "", text)
    text = text.replace("```", "")

    text = re.sub(
        r"^\s{0,3}#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)

    text = re.sub(
        r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)",
        r"\1",
        text
    )

    text = re.sub(
        r"(?<!_)_(?!\s)(.+?)(?<!\s)_(?!_)",
        r"\1",
        text
    )

    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)

    text = re.sub(
        r"^\s*[-*+]\s+",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(
        r"^\s*[-*_]{3,}\s*$",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(
        r"^\s*>\s?",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(
        "["
        "\U0001F300-\U0001F5FF"
        "\U0001F600-\U0001F64F"
        "\U0001F680-\U0001F6FF"
        "\U0001F700-\U0001F77F"
        "\U0001F780-\U0001F7FF"
        "\U0001F800-\U0001F8FF"
        "\U0001F900-\U0001F9FF"
        "\U0001FA00-\U0001FA6F"
        "\U0001FA70-\U0001FAFF"
        "\U00002702-\U000027B0"
        "\U000024C2-\U0001F251"
        "\U0001F1E0-\U0001F1FF"
        "]+",
        "",
        text,
    )

    text = re.sub(
        r"[\u200B-\u200F\u202A-\u202E\u2060\uFEFF]",
        "",
        text
    )

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201C": '"',
        "\u201D": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u00A0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    ai_starters = [
        r"^sure[,!]?\s*",
        r"^certainly[,!]?\s*",
        r"^of course[,!]?\s*",
        r"^here(?:'s| is)\b.*?:\s*",
        r"^below is\b.*?:\s*",
        r"^the following is\b.*?:\s*",
    ]

    for pattern in ai_starters:
        text = re.sub(
            pattern,
            "",
            text,
            flags=re.IGNORECASE
        )

    return text.strip()


# =========================================================
# HELPER: SPLIT TEXT INTO PARAGRAPHS
# =========================================================
def split_text_paragraphs(text):
    text = clean_text(text)

    if not text:
        return []

    paragraphs = re.split(
        r"\n\s*\n+",
        text
    )

    return [
        p.strip()
        for p in paragraphs
        if p.strip()
    ]


# =========================================================
# HELPER: SPLIT INTO BULLET ITEMS
# =========================================================
def split_bullet_items(text):
    text = clean_text(text)

    if not text:
        return []

    chunks = re.split(
        r"\n\s*\n+",
        text
    )

    items = []

    for chunk in chunks:
        lines = [
            line.strip()
            for line in chunk.split("\n")
            if line.strip()
        ]

        if len(lines) > 1:
            items.extend(lines)

        elif lines:
            items.append(lines[0])

    cleaned = []

    for item in items:

        item = re.sub(
            r"^\s*(?:\(?[0-9a-zA-Z]+\)?[\.\)\:]\s+)+",
            "",
            item
        )

        item = re.sub(
            r"^\s*[-\u2013\u2014]\s+",
            "",
            item
        )

        item = item.strip()

        if item:
            cleaned.append(item)

    return cleaned


# =========================================================
# HELPER: ADD BULLET
# =========================================================
def add_bullet(document, text):
    text = str(text or "").strip()

    if not text:
        return

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.paragraph_format.left_indent = Mm(8)
    paragraph.paragraph_format.first_line_indent = Mm(-4)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(4)

    run = paragraph.add_run(text)

    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    return paragraph


# =========================================================
# GETTERS
# =========================================================
def get_student_name(placement):
    student = placement.student
    name = ""

    try:
        name = student.get_full_name() or ""
    except Exception:
        name = ""

    if not name:
        name = getattr(student, "username", "")

    return str(name or "")


def get_registration_number(placement):
    reg = getattr(
        placement,
        "registration_number",
        ""
    )

    if not reg:
        reg = getattr(
            placement.student,
            "registration_number",
            ""
        )

    return str(reg or "")


def get_organization(placement):
    organization = (
        getattr(placement, "organization", "")
        or getattr(placement, "company_name", "")
        or getattr(placement, "organization_name", "")
    )

    return str(organization or "")


def get_institute_name(placement):
    name = (
        getattr(placement, "institute_name", "")
        or getattr(
            placement.student,
            "institute_name",
            ""
        )
        or getattr(
            placement.student,
            "college_name",
            ""
        )
    )

    return str(name or "")


def get_programme(placement):
    programme = (
        getattr(placement, "programme", "")
        or getattr(placement.student, "programme", "")
        or getattr(placement.student, "course", "")
    )

    return str(programme or "")


def get_institute_supervisor(placement):
    name = (
        getattr(
            placement,
            "institute_supervisor",
            ""
        )
        or getattr(
            placement,
            "academic_supervisor",
            ""
        )
    )

    return str(name or "")


# =========================================================
# HELPER: ADD CLOUDINARY IMAGE
# =========================================================
def add_cloudinary_image(
    document,
    cloudinary_image,
    caption=None,
    width_mm=80
):
    if not cloudinary_image:
        return False

    image_url = (
        getattr(cloudinary_image, "url", "")
        or ""
    )

    if not image_url:
        return False

    try:
        request = Request(
            image_url,
            headers={"User-Agent": "Mozilla/5.0"}
        )

        with urlopen(request, timeout=20) as response:
            image_bytes = response.read()

        if not image_bytes:
            return False

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = None

        run = paragraph.add_run()

        run.add_picture(
            BytesIO(image_bytes),
            width=Mm(width_mm)
        )

        if caption:
            caption_paragraph = document.add_paragraph()

            caption_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            caption_paragraph.paragraph_format.first_line_indent = None

            caption_run = caption_paragraph.add_run(
                str(caption)
            )

            caption_run.italic = True
            caption_run.font.name = "Times New Roman"
            caption_run.font.size = Pt(10)

        return True

    except Exception as error:
        print(
            "CLOUDINARY IMAGE DOWNLOAD ERROR:",
            error
        )

        return False


# =========================================================
# BUILD FIELD REPORT DOCX
# =========================================================
def build_field_report_docx(placement, reports):

    document = Document()

    # Ensure Heading styles use Times New Roman
    for style_name in (
        "Heading 1",
        "Heading 2",
        "Heading 3"
    ):
        try:
            style = document.styles[style_name]

            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(
                0x00,
                0x00,
                0x00
            )

        except KeyError:
            pass

    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman"
    )

    # Enable auto-update of fields
    enable_update_fields(document)

    # -----------------------------------------------------
    # GET SAVED REPORT
    # -----------------------------------------------------
    generated_report = (
        GeneratedFieldReport.objects
        .filter(placement=placement)
        .first()
    )

    # -----------------------------------------------------
    # PULL SAVED TEXT
    # -----------------------------------------------------
    historical_background_text = ""
    organization_structure_text = ""
    vision_text = ""
    mission_text = ""
    objectives_text = ""
    overview_text = ""
    activities_text = ""
    observation_text = ""
    challenges_text = ""
    solutions_text = ""
    conclusion_text = ""
    recommendations_text = ""
    references_text = ""

    if generated_report:

        historical_background_text = clean_text(
            generated_report.historical_background or ""
        )

        organization_structure_text = clean_text(
            generated_report.organization_structure or ""
        )

        vision_text = clean_text(
            generated_report.vision or ""
        )

        mission_text = clean_text(
            generated_report.mission or ""
        )

        objectives_text = clean_text(
            generated_report.objectives or ""
        )

        overview_text = clean_text(
            generated_report.section_21_brief_overview or ""
        )

        activities_text = clean_text(
            generated_report.section_22_activities_performed or ""
        )

        observation_text = clean_text(
            generated_report.section_23_general_observation or ""
        )

        challenges_text = clean_text(
            generated_report.section_24_challenges or ""
        )

        solutions_text = clean_text(
            generated_report.section_25_solutions or ""
        )

        conclusion_text = clean_text(
            generated_report.conclusion or ""
        )

        recommendations_text = clean_text(
            generated_report.recommendations or ""
        )

        references_text = clean_text(
            generated_report.references_bibliography or ""
        )

    # -----------------------------------------------------
    # PLACEMENT DATA
    # -----------------------------------------------------
    student_name = get_student_name(placement)
    registration_number = get_registration_number(placement)
    organization = get_organization(placement)
    institute = get_institute_name(placement)
    programme = get_programme(placement)

    department = str(
        getattr(placement, "department", "")
        or ""
    )

    supervisor_name = str(
        getattr(placement, "supervisor_name", "")
        or ""
    )

    institute_supervisor = get_institute_supervisor(
        placement
    )

    start_date = getattr(
        placement,
        "start_date",
        None
    )

    end_date = getattr(
        placement,
        "end_date",
        None
    )

    # =====================================================
    # COVER PAGE
    # =====================================================
    document.add_paragraph()

    add_center_heading(
        document,
        institute or "UNIVERSITY / COLLEGE NAME",
        size=16,
        bold=True,
    )

    document.add_paragraph()

    logo = getattr(
        placement,
        "logo",
        None
    )

    if logo:
        add_cloudinary_image(
            document,
            logo,
            None,
            width_mm=40
        )

    document.add_paragraph()

    if department:
        add_center_heading(
            document,
            department.upper(),
            size=13,
            bold=True
        )

    else:
        add_center_heading(
            document,
            "FACULTY / DEPARTMENT OF COMPUTING AND INFORMATICS",
            size=12,
            bold=True,
        )

    document.add_paragraph()
    document.add_paragraph()

    add_center_heading(
        document,
        "FIELD PRACTICAL TRAINING REPORT",
        size=18,
        bold=True,
    )

    document.add_paragraph()
    document.add_paragraph()

    add_center_heading(
        document,
        "STUDENT PARTICULARS",
        size=12,
        bold=True
    )

    add_detail_line(
        document,
        "Registration Number:",
        registration_number or "-"
    )

    add_detail_line(
        document,
        "Student Name:",
        student_name or "-"
    )

    add_detail_line(
        document,
        "Programme:",
        programme or "-"
    )

    document.add_paragraph()

    add_center_heading(
        document,
        "INFORMATION ON THE UNDERTAKING OF IPT",
        size=12,
        bold=True,
    )

    period_text = ""

    if start_date and end_date:
        period_text = (
            f"{start_date.strftime('%d %B %Y')} to "
            f"{end_date.strftime('%d %B %Y')}"
        )

    elif start_date:
        period_text = start_date.strftime(
            "%d %B %Y"
        )

    add_detail_line(
        document,
        "Organization:",
        organization or "-"
    )

    add_detail_line(
        document,
        "Training Period:",
        period_text or "-"
    )

    document.add_paragraph()

    add_center_heading(
        document,
        "PARTICULARS OF SUPERVISORS",
        size=12,
        bold=True
    )

    add_detail_line(
        document,
        "Institute Supervisor:",
        institute_supervisor or "-"
    )

    add_detail_line(
        document,
        "Industrial Supervisor:",
        supervisor_name or "-"
    )

    document.add_paragraph()
    document.add_paragraph()

    if end_date:
        add_center_heading(
            document,
            end_date.strftime("%B %Y").upper(),
            size=12,
            bold=True,
        )

    # =====================================================
    # NEW SECTION — REST OF REPORT
    # =====================================================
    document.add_section(WD_SECTION.NEW_PAGE)

    try:
        cover_section = document.sections[0]

        configure_section(
            cover_section,
            with_border=True
        )

    except IndexError:
        pass

    try:
        body_section = document.sections[1]

        configure_section(
            body_section,
            with_border=False
        )

        footer = body_section.footer
        footer.is_linked_to_previous = False

        add_page_number(
            footer.paragraphs[0]
        )

    except IndexError:
        pass

    # =====================================================
    # DECLARATION
    # =====================================================
    add_heading_1(
        document,
        "DECLARATION"
    )

    add_body_paragraph(
        document,
        f"I, {student_name}, hereby declare that this Field Practical "
        f"Training Report is my own work and has been prepared based on "
        f"the practical training activities, experiences and knowledge "
        f"acquired during my field placement at {organization}."
    )

    add_body_paragraph(
        document,
        "The information presented in this report represents my personal "
        "experience, observations, knowledge and skills gained during the "
        "field practical training. It has not been submitted elsewhere for "
        "any academic award."
    )

    document.add_paragraph()
    document.add_paragraph()

    add_normal_paragraph(
        document,
        "Student Name: " + (student_name or "")
    )

    add_normal_paragraph(
        document,
        "Signature: __________________________________________"
    )

    add_normal_paragraph(
        document,
        "Date: _______________________________________________"
    )

    document.add_paragraph()

    add_normal_paragraph(
        document,
        "Industrial Supervisor: "
        + (
            supervisor_name
            or "____________________________"
        )
    )

    add_normal_paragraph(
        document,
        "Signature: __________________________________________"
    )

    add_normal_paragraph(
        document,
        "Date: _______________________________________________"
    )

    add_page_break(document)

    # =====================================================
    # TABLE OF CONTENTS
    # =====================================================
    add_heading_1(
        document,
        "TABLE OF CONTENTS"
    )

    add_toc_field(document)

    add_page_break(document)

    # =====================================================
    # LIST OF TABLES
    # =====================================================
    add_heading_1(
        document,
        "LIST OF TABLES"
    )

    if reports:
        add_normal_paragraph(
            document,
            "Table 1: Activities Performed During Industrial Practical Training"
        )

    else:
        add_normal_paragraph(
            document,
            "No tables are included in this report."
        )

    add_page_break(document)

    # =====================================================
    # LIST OF FIGURES
    # =====================================================
    add_heading_1(
        document,
        "LIST OF FIGURES"
    )

    figure_items = []

    if generated_report:

        if generated_report.historical_background_image:
            figure_items.append(
                "Figure 1.1: Historical Background of the Organization"
            )

        if generated_report.organization_structure_image:
            figure_items.append(
                "Figure 1.2: Organization Structure"
            )

        if generated_report.organization_image:
            figure_items.append(
                "Figure 1.3: Organization"
            )

        if generated_report.activity_image:
            figure_items.append(
                "Figure 2.1: Field Training Activity"
            )

        if generated_report.appendix_image:
            figure_items.append(
                "Figure A.1: Supporting Training Image"
            )

    if figure_items:

        for figure in figure_items:
            add_normal_paragraph(
                document,
                figure
            )

    else:
        add_normal_paragraph(
            document,
            "No figures are included in this report."
        )

    add_page_break(document)

    # =====================================================
    # LIST OF ABBREVIATIONS
    # =====================================================
    add_heading_1(
        document,
        "LIST OF ABBREVIATIONS"
    )

    abbreviations = getattr(
        placement,
        "abbreviations",
        ""
    )

    abbreviation_items = []

    if abbreviations:
        abbreviation_items = split_text_paragraphs(
            abbreviations
        )

    if not abbreviation_items:
        abbreviation_items = [
            "FPT - Field Practical Training",
            "IPT - Industrial Practical Training",
            "ICT - Information and Communication Technology",
        ]

    for item in abbreviation_items:
        add_normal_paragraph(
            document,
            item
        )

    add_page_break(document)

    # =====================================================
    # CHAPTER ONE
    # =====================================================
    add_heading_1(
        document,
        "CHAPTER ONE"
    )

    add_heading_1(
        document,
        "INTRODUCTION"
    )

    # -----------------------------------------------------
    # 1.1
    # -----------------------------------------------------
    add_heading_2(
        document,
        "1.1 Background of the Industry/Organization"
    )

    if historical_background_text:

        for paragraph_text in split_text_paragraphs(
            historical_background_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    else:
        add_body_paragraph(
            document,
            f"{organization} is the organization where the field practical "
            f"training was undertaken. The training provided an opportunity "
            f"for the student to gain practical exposure to the working "
            f"environment and to relate theoretical knowledge acquired in "
            f"class with practical activities."
        )

    if (
        generated_report
        and generated_report.historical_background_image
    ):
        add_cloudinary_image(
            document,
            generated_report.historical_background_image,
            "Figure 1.1: Historical Background of the Organization",
            width_mm=80,
        )

    # -----------------------------------------------------
    # 1.2
    # -----------------------------------------------------
    add_heading_2(
        document,
        "1.2 Organization Structure of Industry/Organization"
    )

    if organization_structure_text:

        for paragraph_text in split_text_paragraphs(
            organization_structure_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    else:
        add_body_paragraph(
            document,
            "The organization operates through different departments and "
            "personnel responsible for carrying out its activities. During "
            "the field training, the student worked under the guidance of "
            "supervisors and other responsible staff members."
        )

    if (
        generated_report
        and generated_report.organization_structure_image
    ):
        add_cloudinary_image(
            document,
            generated_report.organization_structure_image,
            "Figure 1.2: Organization Structure",
            width_mm=80,
        )

    # -----------------------------------------------------
    # 1.3
    # -----------------------------------------------------
    add_heading_2(
        document,
        "1.3 Vision, Mission and Objectives of the Industry/Organization"
    )

    # 1.3.1
    add_heading_3(
        document,
        "1.3.1 Vision"
    )

    if vision_text:

        for paragraph_text in split_text_paragraphs(
            vision_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    else:
        add_body_paragraph(
            document,
            "The vision statement of the organization is presented in this section."
        )

    # 1.3.2
    add_heading_3(
        document,
        "1.3.2 Mission"
    )

    if mission_text:

        for paragraph_text in split_text_paragraphs(
            mission_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    else:
        add_body_paragraph(
            document,
            "The mission statement of the organization is presented in this section."
        )

    # 1.3.3
    add_heading_3(
        document,
        "1.3.3 Objectives"
    )

    if objectives_text:

        objective_items = split_bullet_items(
            objectives_text
        )

        if objective_items:

            for obj in objective_items:
                add_bullet(
                    document,
                    obj
                )

        else:
            add_body_paragraph(
                document,
                objectives_text
            )

    else:
        add_body_paragraph(
            document,
            "The objectives of the organization are presented in this section."
        )

    if (
        generated_report
        and generated_report.organization_image
    ):
        add_cloudinary_image(
            document,
            generated_report.organization_image,
            "Figure 1.3: Organization",
            width_mm=80,
        )

    add_page_break(document)

    # =====================================================
    # CHAPTER TWO
    # =====================================================
    add_heading_1(
        document,
        "CHAPTER TWO"
    )

    add_heading_1(
        document,
        "ACTIVITIES PERFORMED"
    )

    # -----------------------------------------------------
    # 2.1
    # -----------------------------------------------------
    add_heading_2(
        document,
        "2.1 Brief Overview of the Undertaking"
    )

    if overview_text:

        for paragraph_text in split_text_paragraphs(
            overview_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    # -----------------------------------------------------
    # 2.2
    # -----------------------------------------------------
    add_heading_2(
        document,
        "2.2 Activity/Activities Performed During IPT"
    )

    if activities_text:

        for paragraph_text in split_text_paragraphs(
            activities_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    # -----------------------------------------------------
    # 2.3
    # -----------------------------------------------------
    add_heading_2(
        document,
        "2.3 General Observations"
    )

    if observation_text:

        for paragraph_text in split_text_paragraphs(
            observation_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    # -----------------------------------------------------
    # 2.4
    # CHANGED: DISPLAY AS NORMAL PARAGRAPHS
    # =====================================================
    add_heading_2(
        document,
        "2.4 Challenges/Problems Faced if Any"
    )

    if challenges_text:

        for paragraph_text in split_text_paragraphs(
            challenges_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    # -----------------------------------------------------
    # 2.5
    # CHANGED: DISPLAY AS NORMAL PARAGRAPHS
    # =====================================================
    add_heading_2(
        document,
        "2.5 How the Student Solved the Challenges"
    )

    if solutions_text:

        for paragraph_text in split_text_paragraphs(
            solutions_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    if (
        generated_report
        and generated_report.activity_image
    ):
        add_cloudinary_image(
            document,
            generated_report.activity_image,
            "Figure 2.1: Field Training Activity",
            width_mm=80,
        )

    add_page_break(document)

    # =====================================================
    # CHAPTER THREE
    # =====================================================
    add_heading_1(
        document,
        "CHAPTER THREE"
    )

    add_heading_1(
        document,
        "CONCLUSION AND RECOMMENDATIONS"
    )

    # -----------------------------------------------------
    # 3.1
    # -----------------------------------------------------
    add_heading_2(
        document,
        "3.1 Conclusion"
    )

    if conclusion_text:

        for paragraph_text in split_text_paragraphs(
            conclusion_text
        ):
            add_body_paragraph(
                document,
                paragraph_text
            )

    # -----------------------------------------------------
    # 3.2
    # -----------------------------------------------------
    add_heading_2(
        document,
        "3.2 Recommendations"
    )

    if recommendations_text:

        recommendation_items = split_bullet_items(
            recommendations_text
        )

        if len(recommendation_items) > 1:

            for rec in recommendation_items:
                add_bullet(
                    document,
                    rec
                )

        else:

            for paragraph_text in split_text_paragraphs(
                recommendations_text
            ):
                add_body_paragraph(
                    document,
                    paragraph_text
                )

    add_page_break(document)

    # =====================================================
    # REFERENCES
    # =====================================================
    add_heading_1(
        document,
        "REFERENCES"
    )

    if references_text:

        reference_items = re.split(
            r"\n\s*\n+|\n(?=\d+\.)|(?<=\.)\s+(?=[A-Z][a-z]+,)",
            references_text,
        )

        for reference in reference_items:

            reference = reference.strip()

            if not reference:
                continue

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.first_line_indent = Mm(0)
            paragraph.paragraph_format.left_indent = Mm(0)

            run = paragraph.add_run(
                reference
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    add_page_break(document)

    # =====================================================
    # APPENDICES
    # =====================================================
    add_heading_1(
        document,
        "APPENDICES"
    )

    add_heading_2(
        document,
        "Appendix A: Daily Field Training Reports"
    )

    if reports:

        table = document.add_table(
            rows=1,
            cols=4
        )

        table.style = "Table Grid"
        table.autofit = False

        header = table.rows[0]

        set_repeat_table_header(
            header
        )

        headers = [
            "No.",
            "Date",
            "Activity",
            "Supervisor Comment"
        ]

        widths = [
            15,
            30,
            90,
            45
        ]

        for index, header_text in enumerate(
            headers
        ):

            cell = header.cells[index]

            set_cell_text(
                cell,
                header_text,
                bold=True
            )

            set_cell_shading(
                cell,
                "D9EAD3"
            )

            set_cell_width(
                cell,
                widths[index]
            )

        for index, report in enumerate(
            reports,
            start=1
        ):

            row = table.add_row()

            values = [
                str(index),
                (
                    report.date.strftime("%d/%m/%Y")
                    if report.date
                    else ""
                ),
                (
                    report.title
                    or report.activities
                    or ""
                ),
                report.supervisor_comment or "",
            ]

            for cell_index, value in enumerate(
                values
            ):

                cell = row.cells[cell_index]

                set_cell_text(
                    cell,
                    value
                )

                set_cell_width(
                    cell,
                    widths[cell_index]
                )

    add_page_break(document)

    add_heading_2(
        document,
        "Appendix B: Supporting Training Documents"
    )

    if (
        generated_report
        and generated_report.appendix_image
    ):
        add_cloudinary_image(
            document,
            generated_report.appendix_image,
            "Figure A.1: Supporting Training Image",
            width_mm=80,
        )

    return document


# =========================================================
# DOWNLOAD DOCX
# =========================================================
@login_required(login_url="home")
@require_POST
def download_field_report_docx(request, id):

    placement = get_object_or_404(
        FieldPlacement,
        id=id
    )

    if placement.student != request.user:
        return JsonResponse(
            {
                "success": False,
                "message": "You are not authorized to download this report.",
            },
            status=403,
        )

    # -----------------------------------------------------
    # COMPLETENESS CHECK
    # -----------------------------------------------------
    generated_report = (
        GeneratedFieldReport.objects
        .filter(placement=placement)
        .first()
    )

    missing = get_missing_report_fields(
        generated_report
    )

    if missing:

        missing_labels = [
            label
            for _, label in missing
        ]

        return JsonResponse(
            {
                "success": False,
                "message": (
                    "The report is not yet complete. "
                    "Please save the following sections before downloading:\n"
                    + "\n".join(
                        f"- {label}"
                        for label in missing_labels
                    )
                ),
                "missing_fields": [
                    field
                    for field, _ in missing
                ],
                "missing_labels": missing_labels,
            },
            status=400,
        )

    # -----------------------------------------------------
    # GET DAILY REPORTS
    # -----------------------------------------------------
    reports = (
        placement.daily_reports
        .all()
        .order_by("date", "id")
    )

    # -----------------------------------------------------
    # BUILD DOCX
    # -----------------------------------------------------
    try:

        document = build_field_report_docx(
            placement,
            reports
        )

        buffer = BytesIO()

        document.save(buffer)

        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        filename = (
            f"Field_Practical_Training_Report_"
            f"{get_registration_number(placement) or placement.id}.docx"
        )

        response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )

        return response

    except Exception as e:

        import traceback

        traceback.print_exc()

        return JsonResponse(
            {
                "success": False,
                "message": "Failed to generate DOCX report.",
                "error": str(e),
            },
            status=500,
        )